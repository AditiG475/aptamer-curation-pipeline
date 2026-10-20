"""Stage 6 - Export a formatted, filterable Excel workbook.

Produces a review-ready spreadsheet: frozen header, AutoFilter dropdowns,
colour-coded ligand roles, bold high-impact journals, clickable DOI links
and a second sheet summarising data-quality flags.

Usage:
    python pipeline/s06_export_excel.py --input out/05_structures.csv --output out/curated_dataset.xlsx
"""
import argparse

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROLE_FILL = {
    "Biological target": "D9EAD3",
    "Buffer / salt component": "FCE5CD",
    "Crystallographic additive": "D0E0F2",
}
STATUS_COLUMNS = ["kd_status", "doi_status", "pubchem_status", "structure_status"]


def quality_summary(df):
    rows = []
    for column in STATUS_COLUMNS:
        if column in df:
            for status, count in df[column].value_counts(dropna=False).items():
                rows.append({"check": column, "status": status, "entries": count})
    return pd.DataFrame(rows)


def format_sheet(ws, df):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3C4A5C")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    columns = {name: idx + 1 for idx, name in enumerate(df.columns)}
    for row in range(2, ws.max_row + 1):
        role = ws.cell(row, columns["ligand_role"]).value if "ligand_role" in columns else None
        if role in ROLE_FILL:
            ws.cell(row, columns["ligand_role"]).fill = PatternFill("solid", fgColor=ROLE_FILL[role])
        if "journal_tier" in columns and ws.cell(row, columns["journal_tier"]).value == "High-impact":
            ws.cell(row, columns["journal"]).font = Font(bold=True)
        if "doi" in columns:
            doi_cell = ws.cell(row, columns["doi"])
            if doi_cell.value:
                doi_cell.hyperlink = f"https://doi.org/{doi_cell.value}"
                doi_cell.font = Font(color="0563C1", underline="single")

    for name, idx in columns.items():
        longest = max([len(str(name))] + [len(str(v)) for v in df[name].head(200)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    df = pd.read_csv(args.input)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="dataset", index=False)
        quality_summary(df).to_excel(writer, sheet_name="quality_flags", index=False)

    wb = load_workbook(args.output)
    format_sheet(wb["dataset"], df)
    wb.save(args.output)
    print(f"[stage 6] Wrote {len(df)} entries to {args.output}")


if __name__ == "__main__":
    main()
